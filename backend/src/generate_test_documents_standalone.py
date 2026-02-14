"""
Script para generar documentos de prueba sin necesidad de base de datos

Este script crea archivos de ejemplo en las rutas correctas para que
los usuarios de prueba puedan visualizar sus documentos en el dashboard.
"""

import os
import uuid
from datetime import datetime

# IDs de usuario de prueba (estos son los IDs que se generan en seed.py)
TEST_USERS = [
    {
        "email": "admin@repa.gob.ar",
        "user_id": "4938043f-6619-4e55-b546-e40857b3b9a6",
        "nombre": "María",
        "apellido": "González",
        "tipo": "pf",  # Persona Física
    },
    {
        "email": "usuario1@repa.gob.ar",
        "user_id": "12345678-1234-1234-1234-123456789012",
        "nombre": "Juan",
        "apellido": "Pérez",
        "razon_social": "Productora Misionera SRL",
        "tipo": "pj",  # Persona Jurídica
    },
    {
        "email": "usuario2@repa.gob.ar",
        "user_id": "87654321-4321-4321-4321-210987654321",
        "nombre": "Luciana",
        "apellido": "Fernández",
        "tipo": "pf",
    },
    {
        "email": "usuario3@repa.gob.ar",
        "user_id": "11111111-1111-1111-1111-111111111111",
        "nombre": "Pedro",
        "apellido": "Martínez",
        "razon_social": "Colectivo Audiovisual del NEA",
        "tipo": "pj",
    },
]

# Directorio base para uploads
if os.name == "nt":  # Windows
    UPLOAD_BASE_DIR = os.path.join(
        os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
        "uploads",
    )
else:
    UPLOAD_BASE_DIR = "/app/uploads"


def get_user_upload_dir(user_id: str) -> str:
    """Obtener el directorio de uploads para un usuario específico"""
    return os.path.join(UPLOAD_BASE_DIR, user_id)


def create_sample_file(filepath: str, title: str, extension: str):
    """Crear un archivo de ejemplo simple"""
    # Asegurar que el directorio exista
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    if extension.lower() == ".pdf":
        try:
            # Intentar crear PDF si reportlab está disponible
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.units import inch
            from reportlab.pdfgen import canvas

            c = canvas.Canvas(filepath, pagesize=letter)
            width, height = letter

            c.setFont("Helvetica-Bold", 16)
            c.drawString(inch, height - inch, title)

            c.setFont("Helvetica", 12)
            y_position = height - 2 * inch

            c.drawString(
                inch, y_position, "Documento de ejemplo generado automáticamente"
            )
            y_position -= 0.5 * inch
            c.drawString(
                inch,
                y_position,
                f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
            )
            y_position -= 0.5 * inch
            c.drawString(inch, y_position, f"ID único: {uuid.uuid4().hex[:8]}")
            y_position -= 0.5 * inch
            c.drawString(
                inch, y_position, "Este es un documento de prueba para el sistema RePA."
            )

            c.save()
            print(f"  ✓ PDF creado: {os.path.basename(filepath)}")
            return True
        except ImportError:
            # Si no hay reportlab, crear un archivo de texto con extensión .txt
            filepath = filepath.replace(".pdf", ".txt")

    if extension.lower() in [".docx", ".txt"] or filepath.endswith(".txt"):
        try:
            # Intentar crear DOCX si python-docx está disponible
            if extension.lower() == ".docx" and not filepath.endswith(".txt"):
                from docx import Document

                doc = Document()
                doc.add_heading(title, 0)

                p = doc.add_paragraph(
                    "Documento de ejemplo generado automáticamente\n\n"
                )
                p.add_run(
                    f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n"
                )
                p.add_run(f"ID único: {uuid.uuid4().hex[:8]}\n\n")
                p.add_run("Este es un documento de prueba para el sistema RePA.")

                doc.save(filepath)
                print(f"  ✓ DOCX creado: {os.path.basename(filepath)}")
                return True
        except ImportError:
            # Si no hay python-docx, crear un archivo de texto
            filepath = filepath.replace(".docx", ".txt")

    # Crear archivo de texto como fallback
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(f"{title}\n")
        f.write("=" * len(title) + "\n\n")
        f.write("Documento de ejemplo generado automáticamente\n\n")
        f.write(
            f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n"
        )
        f.write(f"ID único: {uuid.uuid4().hex[:8]}\n\n")
        f.write("Este es un documento de prueba para el sistema RePA.\n")
        f.write(f"\nFormato original solicitado: {extension}")

    print(f"  ✓ TXT creado: {os.path.basename(filepath)}")
    return True


def generate_documents():
    """Generar documentos para todos los usuarios de prueba"""
    print("🚀 Iniciando generación de documentos de prueba...")
    print(f"Directorio de uploads: {UPLOAD_BASE_DIR}")

    # Crear directorio base si no existe
    os.makedirs(UPLOAD_BASE_DIR, exist_ok=True)

    for user in TEST_USERS:
        print(f"\n📁 Procesando usuario: {user['email']} ({user['tipo'].upper()})")
        user_dir = get_user_upload_dir(user["user_id"])

        if user["tipo"] == "pf":
            # Persona Física - Solo DNI
            filename = (
                f"DNI_{user['apellido']}_{user['nombre']}_{uuid.uuid4().hex[:8]}.pdf"
            )
            filepath = os.path.join(user_dir, filename)
            create_sample_file(
                filepath, f"DNI - {user['nombre']} {user['apellido']}", ".pdf"
            )

            print(f"  📄 Path para BD: {filename}")

        elif user["tipo"] == "pj":
            # Persona Jurídica - Múltiples documentos
            docs = [
                ("estatuto", f"Estatuto - {user['razon_social']}", ".pdf"),
                (
                    "constancia_cuit",
                    f"Constancia CUIT - {user['razon_social']}",
                    ".pdf",
                ),
                (
                    "acta_autoridades",
                    f"Acta de Autoridades - {user['razon_social']}",
                    ".docx",
                ),
                (
                    "cv_institucional",
                    f"CV Institucional - {user['razon_social']}",
                    ".docx",
                ),
            ]

            print("  📄 Paths para BD:")
            for doc_type, title, ext in docs:
                filename = f"{doc_type}_{user['razon_social'].replace(' ', '_')}_{uuid.uuid4().hex[:8]}{ext}"
                filepath = os.path.join(user_dir, filename)
                create_sample_file(filepath, title, ext)
                print(f"    - {doc_type}: {filename}")

    print("\n✅ Documentos generados exitosamente!")
    print("\n📝 Para actualizar la base de datos con estos paths, ejecuta:")
    print(
        "   UPDATE personas_fisicas SET dni_adjunto_path = '<filename>' WHERE user_id = '<user_id>';"
    )
    print(
        "   UPDATE persona_juridica SET estatuto_path = '<filename>' WHERE user_id = '<user_id>';"
    )
    print("   ... y así con los demás campos")


if __name__ == "__main__":
    generate_documents()
