"""
Script para generar documentos de prueba para los usuarios creados en seed.py

Este script crea archivos PDF de ejemplo en las rutas correctas para que
los usuarios de prueba puedan visualizar sus documentos en el dashboard.
"""

import os
import uuid
from datetime import datetime

from src.database import SessionLocal
from src.models.persona_fisica_model import PersonaFisica
from src.models.persona_juridica_model import PersonaJuridica
from src.models.user_models import User

# Directorio base para uploads (debe coincidir con upload_routes.py)
UPLOAD_BASE_DIR = "/app/uploads"


def get_user_upload_dir(user_id: str) -> str:
    """Obtener el directorio de uploads para un usuario específico"""
    return os.path.join(UPLOAD_BASE_DIR, user_id)


def create_sample_pdf(filepath: str, title: str):
    """Crear un PDF de ejemplo simple"""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas

    # Asegurar que el directorio exista
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    c = canvas.Canvas(filepath, pagesize=letter)
    width, height = letter

    # Título
    c.setFont("Helvetica-Bold", 16)
    c.drawString(inch, height - inch, title)

    # Contenido de ejemplo
    c.setFont("Helvetica", 12)
    y_position = height - 2 * inch

    c.drawString(inch, y_position, "Documento de ejemplo generado automáticamente")
    y_position -= 0.5 * inch
    c.drawString(
        inch,
        y_position,
        f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
    )
    y_position -= 0.5 * inch
    c.drawString(inch, y_position, f"ID único: {uuid.uuid4().hex[:8]}")
    y_position -= 0.5 * inch
    c.drawString(
        inch, y_position, "Este es un documento de prueba para el sistema RePA."
    )

    c.save()


def create_sample_docx(filepath: str, title: str):
    """Crear un DOCX de ejemplo simple"""
    from docx import Document

    # Asegurar que el directorio exista
    os.makedirs(os.path.dirname(filepath), exist_ok=True)

    doc = Document()
    doc.add_heading(title, 0)

    p = doc.add_paragraph("Documento de ejemplo generado automáticamente\n\n")
    p.add_run(f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}\n")
    p.add_run(f"ID único: {uuid.uuid4().hex[:8]}\n\n")
    p.add_run("Este es un documento de prueba para el sistema RePA.")

    doc.save(filepath)


def generate_documents():
    """Generar documentos para todos los usuarios que tengan registros"""
    db = SessionLocal()

    try:
        # Generar DNI para Persona Física
        print("Generando DNIs para Persona Física...")
        personas_fisicas = (
            db.query(PersonaFisica)
            .filter(PersonaFisica.dni_adjunto_path.is_(None))
            .all()
        )

        for pf in personas_fisicas:
            user = db.query(User).filter(User.id == pf.user_id).first()
            if user:
                filename = f"DNI_{pf.apellido}_{pf.nombre}_{uuid.uuid4().hex[:8]}.pdf"
                filepath = os.path.join(get_user_upload_dir(pf.user_id), filename)

                # Crear el PDF
                create_sample_pdf(filepath, f"DNI - {pf.nombre} {pf.apellido}")

                # Actualizar la base de datos
                pf.dni_adjunto_path = filename
                db.commit()

                print(f"  ✓ DNI creado para {user.email}: {filename}")

        # Generar documentos para Persona Jurídica
        print("\nGenerando documentos para Persona Jurídica...")
        personas_juridicas = (
            db.query(PersonaJuridica)
            .filter(PersonaJuridica.estatuto_path.is_(None))
            .all()
        )

        for pj in personas_juridicas:
            user = db.query(User).filter(User.id == pj.user_id).first()
            if user:
                # Estatuto
                filename = f"estatuto_{pj.razon_social}_{uuid.uuid4().hex[:8]}.pdf"
                filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
                create_sample_pdf(filepath, f"Estatuto - {pj.razon_social}")
                pj.estatuto_path = filename

                # Constancia CUIT
                filename = (
                    f"constancia_cuit_{pj.razon_social}_{uuid.uuid4().hex[:8]}.pdf"
                )
                filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
                create_sample_pdf(filepath, f"Constancia CUIT - {pj.razon_social}")
                pj.constancia_cuit_path = filename

                # Acta de Autoridades
                filename = (
                    f"acta_autoridades_{pj.razon_social}_{uuid.uuid4().hex[:8]}.docx"
                )
                filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
                create_sample_docx(filepath, f"Acta de Autoridades - {pj.razon_social}")
                pj.acta_autoridades_path = filename

                # CV Institucional
                filename = (
                    f"cv_institucional_{pj.razon_social}_{uuid.uuid4().hex[:8]}.docx"
                )
                filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
                create_sample_docx(filepath, f"CV Institucional - {pj.razon_social}")
                pj.cv_institucional_path = filename

                db.commit()

                print(f"  ✓ Documentos creados para {user.email}:")
                print(f"    - Estatuto: {pj.estatuto_path}")
                print(f"    - Constancia CUIT: {pj.constancia_cuit_path}")
                print(f"    - Acta de Autoridades: {pj.acta_autoridades_path}")
                print(f"    - CV Institucional: {pj.cv_institucional_path}")

        print("\n✅ Documentos generados exitosamente!")

    except ImportError as e:
        print("\n❌ Error: Falta instalar dependencias para generar documentos:")
        print("   pip install reportlab python-docx")
        print(f"\nError detallado: {e}")
    except Exception as e:
        print(f"\n❌ Error al generar documentos: {e}")
        db.rollback()
    finally:
        db.close()


if __name__ == "__main__":
    # Verificar que no estemos en producción
    from src.config import IS_PRODUCTION

    if IS_PRODUCTION:
        print("❌ Este script no debe ejecutarse en producción")
        exit(1)

    print("🚀 Iniciando generación de documentos de prueba...")
    print(f"Directorio de uploads: {UPLOAD_BASE_DIR}")

    # Crear directorio base si no existe
    os.makedirs(UPLOAD_BASE_DIR, exist_ok=True)

    generate_documents()
