"""
Función para generar documentos de prueba - se integra con seed.py
"""
import os
import uuid
from datetime import datetime

# Directorio base para uploads (debe coincidir con upload_routes.py)
if os.name == 'nt':  # Windows
    UPLOAD_BASE_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), "uploads")
else:
    UPLOAD_BASE_DIR = "/app/uploads"

def get_user_upload_dir(user_id: str) -> str:
    """Obtener el directorio de uploads para un usuario específico"""
    return os.path.join(UPLOAD_BASE_DIR, user_id)

def create_sample_pdf(filepath: str, title: str, extra_info: dict = None):
    """Crear un PDF de ejemplo simple"""
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import inch
    
    # Asegurar que el directorio exista
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    
    c = canvas.Canvas(filepath, pagesize=letter)
    width, height = letter
    
    # Título
    c.setFont("Helvetica-Bold", 18)
    c.drawString(inch, height - inch, title)
    
    # Contenido de ejemplo
    c.setFont("Helvetica", 12)
    y_position = height - 2 * inch
    
    # Mostrar información adicional si existe
    if extra_info:
        c.setFont("Helvetica-Bold", 14)
        c.drawString(inch, y_position, "DATOS DEL DOCUMENTO")
        y_position -= 0.4 * inch
        c.setFont("Helvetica", 12)
        for key, value in extra_info.items():
            c.drawString(inch, y_position, f"{key}: {value}")
            y_position -= 0.3 * inch
        y_position -= 0.2 * inch
    
    c.drawString(inch, y_position, "Documento de ejemplo generado automáticamente")
    y_position -= 0.4 * inch
    c.drawString(inch, y_position, f"Fecha de generación: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}")
    y_position -= 0.4 * inch
    doc_id = uuid.uuid4().hex[:8]
    c.drawString(inch, y_position, f"ID único del documento: {doc_id}")
    y_position -= 0.4 * inch
    c.drawString(inch, y_position, "Este es un documento de prueba para el sistema RePA.")
    
    c.save()

def create_sample_docx(filepath: str, title: str):
    """Crear un DOCX de ejemplo simple"""
    from docx import Document
    
    # Asegurar que el directorio exista
    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    
    doc = Document()
    doc.add_heading(title, 0)
    
    p = doc.add_paragraph('Documento de ejemplo generado automáticamente\n\n')
    p.add_run(f'Fecha de generación: {datetime.now().strftime("%d/%m/%Y %H:%M:%S")}\n')
    p.add_run(f'ID único: {uuid.uuid4().hex[:8]}\n\n')
    p.add_run('Este es un documento de prueba para el sistema RePA.')
    
    doc.save(filepath)

def generate_test_documents(db):
    """Generar documentos de prueba para los usuarios creados"""
    from src.models.user_models import User
    from src.models.persona_fisica_model import PersonaFisica
    from src.models.persona_juridica_model import PersonaJuridica
    
    print("\n📁 Generando documentos de prueba...")
    
    # Crear directorio base si no existe
    os.makedirs(UPLOAD_BASE_DIR, exist_ok=True)
    
    # Generar DNI para Persona Física
    personas_fisicas = db.query(PersonaFisica).filter(
        PersonaFisica.dni_adjunto_path.is_(None)
    ).all()
    
    for pf in personas_fisicas:
        user = db.query(User).filter(User.id == pf.user_id).first()
        if user:
            filename = f"DNI_{pf.apellido}_{pf.nombre}_{uuid.uuid4().hex[:8]}.pdf"
            filepath = os.path.join(get_user_upload_dir(pf.user_id), filename)
            
            # Crear el PDF con datos de la persona
            extra_info = {
                "Nombre": pf.nombre or "N/A",
                "Apellido": pf.apellido or "N/A",
                "DNI": pf.dni or "N/A",
                "CUIL": pf.cuil or "N/A",
                "Email": user.email
            }
            create_sample_pdf(filepath, f"DNI - {pf.nombre} {pf.apellido}", extra_info)
            
            # Actualizar la base de datos
            pf.dni_adjunto_path = filename
            print(f"  ✓ DNI creado para {user.email}: {filename}")
    
    # Generar documentos para Persona Jurídica
    personas_juridicas = db.query(PersonaJuridica).filter(
        PersonaJuridica.estatuto_path.is_(None)
    ).all()
    
    for pj in personas_juridicas:
        user = db.query(User).filter(User.id == pj.user_id).first()
        if user:
            # Estatuto
            filename = f"estatuto_{pj.nombre_pj}_{uuid.uuid4().hex[:8]}.pdf"
            filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
            create_sample_pdf(filepath, f"Estatuto - {pj.nombre_pj}")
            pj.estatuto_path = filename
            
            # Constancia CUIT
            filename = f"constancia_cuit_{pj.nombre_pj}_{uuid.uuid4().hex[:8]}.pdf"
            filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
            create_sample_pdf(filepath, f"Constancia CUIT - {pj.nombre_pj}")
            pj.constancia_cuit_path = filename
            
            # Acta de Autoridades
            filename = f"acta_autoridades_{pj.nombre_pj}_{uuid.uuid4().hex[:8]}.docx"
            filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
            create_sample_docx(filepath, f"Acta de Autoridades - {pj.nombre_pj}")
            pj.acta_autoridades_path = filename
            
            # CV Institucional
            filename = f"cv_institucional_{pj.nombre_pj}_{uuid.uuid4().hex[:8]}.docx"
            filepath = os.path.join(get_user_upload_dir(pj.user_id), filename)
            create_sample_docx(filepath, f"CV Institucional - {pj.nombre_pj}")
            pj.cv_institucional_path = filename
            
            print(f"  ✓ Documentos creados para {user.email}")
    
    # Commit de los cambios
    db.commit()
    print("✅ Documentos generados y base de datos actualizada!")
